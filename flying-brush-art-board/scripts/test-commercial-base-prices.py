"""Synthetic fixtures only; no private workbook is authored or changed."""
import importlib.util
from pathlib import Path
import tempfile
import unittest
import openpyxl

spec = importlib.util.spec_from_file_location("prices", Path(__file__).with_name("commercial-base-prices.py"))
prices = importlib.util.module_from_spec(spec)
spec.loader.exec_module(prices)
builder_spec = importlib.util.spec_from_file_location("builder", Path(__file__).with_name("build-item-value-data.py"))
builder = importlib.util.module_from_spec(builder_spec)
builder_spec.loader.exec_module(builder)


def workbook():
    w = openpyxl.Workbook()
    s = w.active
    s.title = "IN_玩法_道具定价"
    for _ in range(4): s.append([])
    s.append([None, "道具名称", "rmb价值", "钻石价值", "金币价值", "道具id"])
    return w, s


def extract(w, items):
    with tempfile.TemporaryDirectory() as folder:
        path = Path(folder) / "source.xlsx"
        w.save(path)
        before = path.read_bytes()
        result = prices.extract(path, items)
        assert before == path.read_bytes(), "source workbook must remain read-only"
        return result


class ValuationTests(unittest.TestCase):
    def test_direct_column_wins_over_merchandise_prices(self):
        w, s = workbook()
        s.append([None, "材料", 3, 30, 600, 1])
        sales = w.create_sheet("日周月礼包表")
        sales.append([None, "礼包", "材料", 20, 480, 1.3])
        result, _, _ = extract(w, [{"id": 1, "name": "材料"}])
        self.assertEqual(result['1']['value'], 30)
        self.assertIn('!D6', result['1']['source'])
        self.assertNotIn('estimate', result['1'])

    def test_zero_missing_error_and_uncached_formula_are_distinct(self):
        w, s = workbook()
        for n, v in enumerate([0, None, '#REF!', '=C9*10', 0.05], 1):
            s.append([None, f"材料{n}", None, v, None, n])
        result, _, _ = extract(w, [{"id": n, "name": f"材料{n}"} for n in range(1, 6)])
        self.assertEqual([result[str(n)]['value'] for n in range(1, 6)], [0, None, None, None, 0.05])

    def test_duplicate_id_conflict_is_not_averaged(self):
        w, s = workbook()
        for v in [30, 40]: s.append([None, "材料", None, v, None, 1])
        result, _, audit = extract(w, [{"id": 1, "name": "材料"}])
        self.assertIsNone(result['1']['value'])
        self.assertEqual(len(audit['unmatched']), 1)

    def test_id_mapping_distinguishes_fragments_and_preserves_name_difference(self):
        w, s = workbook()
        s.append([None, "角色碎片", None, 5, None, 1])
        s.append([None, "角色", None, 50, None, 2])
        s.append([None, "旧名称", None, 20, None, 3])
        result, _, audit = extract(w, [{"id": 1, "name": "角色", "debris": 1}, {"id": 2, "name": "角色"}, {"id": 3, "name": "新名称"}])
        self.assertEqual(result['1']['displayName'], '角色碎片')
        self.assertEqual(result['2']['value'], 50)
        self.assertEqual(audit['nameDifferences'][0]['id'], 3)

    def test_no_fallback_to_other_id_or_generic_quality(self):
        w, s = workbook()
        s.append([None, "材料", None, 30, None, 99])
        s.append([None, "随机精良宝石", None, 40, None, 50])
        result, _, _ = extract(w, [{"id": 1, "name": "材料"}, {"id": 2, "name": "精良头盔宝石", "quality": 2}])
        self.assertFalse(result)

    def test_wrong_workbook_fails_instead_of_inventing_prices(self):
        w = openpyxl.Workbook()
        with self.assertRaises(ValueError): prices.read_rows(w)

    def test_related_random_items_become_questions_not_invented_prices(self):
        result = [{"id": n, "name": f"部位{n}", "basePrice": None, "description": "", "priceNote": "缺直接定价"} for n in [1, 2]]
        items = {n: {"id": n, "name": f"部位{n}", "item_type": 22, "quality": 2} for n in [1, 2]}
        items[3] = {"id": 3, "name": "随机测试宝石", "item_type": 5, "quality": 2, "use_value": 1, "use_detail_value": "10"}
        drops = [{"id": 10, "weight_drop_award": [{"id": 1, "num": 1}, {"id": 2, "num": 1}]}]
        reviews = builder.pricing_reviews(result, items, drops, {"3": {"value": 42, "source": "Sheet!D6"}}, [])
        self.assertEqual(len(reviews), 1)
        self.assertEqual(reviews[0]['itemIds'], [1, 2])
        self.assertIn('42', reviews[0]['question'])
        self.assertTrue(all(i['basePrice'] is None for i in result))
        w, s = workbook(); s['D5'] = '商品售价'
        with self.assertRaises(ValueError): prices.read_rows(w)


if __name__ == '__main__': unittest.main()
