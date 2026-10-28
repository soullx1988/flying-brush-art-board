"""Read-only extraction of item valuation evidence; never infer from sale price alone.

The supplied workbook has single-item price/rebate relationships rather than a
base-price column. Preserve the spread of all row estimates, not false precision.
Generated values are an exchange artifact; the workbook remains the owner.
"""
from collections import defaultdict
from decimal import Decimal
from statistics import median
import hashlib
import math
from pathlib import Path


def positive(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value) and value > 0


def read_rows(workbook):
    rows = defaultdict(dict)
    skipped = []
    for sheet in ("日周月礼包表", "日周月礼包"):
        if sheet not in workbook.sheetnames:
            continue
        ws = workbook[sheet]
        expected = {3: "道具", 4: "道具数量", 5: "钻石价格", 6: "返利比", 7: "折扣"}
        if any(ws.cell(3, col).value != label for col, label in expected.items()):
            raise ValueError(f"{sheet}: pricing headers changed; review the source before importing")
        for row in ws.iter_rows(min_row=4, max_col=7):
            _, _, item, quantity, price, rebate, discount = row
            if not isinstance(item.value, str) or not item.value.strip():
                continue
            name = item.value.strip()
            where = f"{sheet}!C{item.row}:G{item.row}"
            q, p, f, d = quantity.value, price.value, rebate.value, discount.value
            if not positive(q) or not positive(p):
                skipped.append({"source": where, "name": name, "reason": "免费、广告或缺少有效数量/钻石价格，不能用于基础价"})
                continue
            # Two simultaneous definitions require human review, not an arbitrary preference.
            if positive(f) and positive(d):
                skipped.append({"source": where, "name": name, "reason": "返利比与折扣同时存在，口径待核验"})
                continue
            dec = lambda x: Decimal(str(x))
            if positive(f):
                value = dec(p) * dec(f) / dec(q)
                formula = f"{price.coordinate} × {rebate.coordinate} ÷ {quantity.coordinate} = {p} × {f} ÷ {q}"
            elif positive(d) and d <= 1:
                value = dec(p) / dec(d) / dec(q)
                formula = f"{price.coordinate} ÷ {discount.coordinate} ÷ {quantity.coordinate} = {p} ÷ {d} ÷ {q}"
            else:
                skipped.append({"source": where, "name": name, "reason": "缺有效返利比/折扣；售价不能替代基础价"})
                continue
            key = (q, p, f, d)
            evidence = rows[name].setdefault(key, {"value": float(value), "formula": formula, "sources": []})
            evidence["sources"].append(where)
    return {name: list(values.values()) for name, values in rows.items()}, skipped


def extract(path, items):
    import openpyxl
    path = Path(path)
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        groups, skipped = read_rows(workbook)
    finally:
        workbook.close()
    by_name = defaultdict(list)
    for item in items:
        by_name[item["name"]].append(item["id"])
    prices, unmatched = {}, []
    for name, evidence in groups.items():
        matches = by_name.get(name, [])
        if len(matches) != 1:
            unmatched.append({"name": name, "reason": "当前道具名称未唯一匹配，未自动转译旧名称", "evidence": evidence})
            continue
        values = [e["value"] for e in evidence]
        low, high = min(values), max(values)
        # Large disagreements are not a usable baseline, even as an approximate plot.
        if high / low > 1.15:
            unmatched.append({"name": name, "reason": "单品还原值相差超过 15%，基础价值冲突待核验", "evidence": evidence})
            continue
        sources = list(dict.fromkeys(s for e in evidence for s in e["sources"]))
        prices[str(matches[0])] = {
            "value": median(values),
            "estimate": {"min": low, "max": high},
            "source": f"{path} / " + "、".join(sources),
            "note": "表内推算，非独立基础价列。单品价值 = 钻石价格 × 返利比 ÷ 数量（仅填折扣时为价格 ÷ 折扣 ÷ 数量）。返利比精度有限，保留各行还原值范围，未强行归为一个精确数值。\n"
                + "\n".join(f"{e['sources'][0]}：{e['formula']} → {e['value']:g} 钻石/个" for e in evidence)
                + "\n链式礼包的价值总和未逐项拆分，不参与单价还原；本值不能证明现行线上商品价格。",
        }
    # This is a dimensional identity, not a commercial exchange-rate assumption.
    if len(by_name.get("钻石", [])) == 1:
        prices[str(by_name["钻石"][0])] = {"value": 1, "source": "坐标单位：每 1 钻石以 1 钻石计", "note": "仅作钻石价值单位基准；不是人民币汇率。"}
    provenance = {"path": str(path), "sha256": hashlib.sha256(path.read_bytes()).hexdigest(), "label": "用户指定《商业化产出》：单品价值关系，只读推算"}
    audit = {"source": provenance, "unmatched": unmatched, "skipped": skipped,
             "excluded": ["链式礼包 F 列为整段/礼包价值总和，未逐项归因，且存在 #REF! 公式；不分摊为道具单价。", "未将随机/全职业经验自动映射为英雄经验，未将宠物相关旧名自动映射为魔灵道具。"]}
    return prices, provenance, audit
